"""Excel import templates and tolerant parsers for 2D nesting and 1D cutting jobs.

Parsers never raise on bad rows: they return row-level Issues (ref "Sheet!row N") and skip the row.
Header matching ignores case, spaces, underscores and hyphens.
"""
from __future__ import annotations

import io
import re
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .models import Issue, Piece1D, Piece2D, Sheet2D, Stock1D

NOTE = ("DEMO template - ITMC FenestrationOS. Lengths in mm. Blank Qty in Stock = unlimited. "
        "Y/N columns accept Y, N, Yes, No, True, False, 1, 0.")

NEST_PIECES = ["Mark", "Width_mm", "Height_mm", "Qty", "Material", "Thickness_mm", "CanRotate"]
NEST_STOCK = ["Id", "Width_mm", "Height_mm", "Qty", "Material", "Thickness_mm", "IsResidual", "CostPerSheet"]
NEST_SETTINGS = [("Kerf_mm", 3.0, "Minimum gap between pieces"),
                 ("TrimLeft_mm", 0.0, "Edge trim"), ("TrimRight_mm", 0.0, "Edge trim"),
                 ("TrimTop_mm", 0.0, "Edge trim"), ("TrimBottom_mm", 0.0, "Edge trim"),
                 ("Guillotine", "Y", "Y = edge-to-edge cuts only (glass)"),
                 ("MinResidualWidth_mm", 300.0, "Minimum reusable residual width"),
                 ("MinResidualHeight_mm", 300.0, "Minimum reusable residual height"),
                 ("PreferResiduals", "Y", "Use residual sheets first")]
CUT_PIECES = ["Mark", "Length_mm", "Qty", "AngleStart", "AngleEnd", "Label"]
CUT_STOCK = ["Id", "Length_mm", "Qty", "IsOffcut", "CostPerBar"]
CUT_SETTINGS = [("Kerf_mm", 4.0, "Saw kerf per cut"), ("TrimStart_mm", 10.0, "Squaring cut at bar start"),
                ("TrimEnd_mm", 0.0, "Unusable clamp zone at bar end"),
                ("MinReusableOffcut_mm", 500.0, "Remainders >= this are reusable"),
                ("Material", "", "Profile code / description"),
                ("PreferOffcuts", "Y", "Use offcut stock first")]


def _key(s: Any) -> str:
    return re.sub(r"[\s_\-()]+", "", str(s or "")).lower()


def _wb_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet(wb: Workbook, title: str, headers: List[str], rows: List[List[Any]], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE6F0")
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, len(h) + 4)
    return ws


def _settings_sheet(wb: Workbook, settings: List[Tuple[str, Any, str]]):
    ws = _sheet(wb, "Settings", ["Setting", "Value", "Description"], [list(s) for s in settings])
    ws.append([])
    ws.append([NOTE])
    return ws


def build_nesting_template() -> bytes:
    wb = Workbook()
    _sheet(wb, "Pieces", NEST_PIECES, [["G-01", 1200, 2400, 4, "DGU-24", 24, "Y"],
                                       ["P-01", 800, 600, 2, "ACP-4", 4, "N"]], first=True)
    _sheet(wb, "Stock", NEST_STOCK, [["SHEET-3210x2250", 3210, 2250, "", "DGU-24", 24, "N", ""],
                                     ["ACP-1500x4000", 1500, 4000, 20, "ACP-4", 4, "N", ""]])
    _settings_sheet(wb, NEST_SETTINGS)
    return _wb_bytes(wb)


def build_cutting_template() -> bytes:
    wb = Workbook()
    _sheet(wb, "Pieces", CUT_PIECES, [["M1", 3600, 6, 90, 90, "Mullion"], ["T1", 1450, 12, 90, 90, "Transom"]],
           first=True)
    _sheet(wb, "Stock", CUT_STOCK, [["BAR-6000", 6000, "", "N", ""], ["BAR-6500", 6500, "", "N", ""]])
    _settings_sheet(wb, CUT_SETTINGS)
    return _wb_bytes(wb)


# ---------------------------------------------------------------------------
# parsing helpers
# ---------------------------------------------------------------------------
def _find_ws(wb, name: str):
    for ws in wb.worksheets:
        if _key(ws.title) == _key(name):
            return ws
    return None


def _rows(ws, sheet_name: str, issues: List[Issue]) -> List[Tuple[int, Dict[str, Any]]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [_key(h) for h in rows[0]]
    out = []
    for n, r in enumerate(rows[1:], start=2):
        if r is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in r):
            continue
        d = {header[i]: v for i, v in enumerate(r) if i < len(header) and header[i]}
        out.append((n, d))
    return out


def _num(d: Dict[str, Any], names: List[str], *, required: bool, positive: bool, integer: bool,
         issues: List[Issue], ref: str, label: str, default=None, allow_zero: bool = False):
    v = None
    for nm in names:
        if _key(nm) in d and d[_key(nm)] is not None and str(d[_key(nm)]).strip() != "":
            v = d[_key(nm)]
            break
    if v is None:
        if required:
            issues.append(Issue(code="MISSING_VALUE", severity="error", message=f"{label} is missing", ref=ref))
            return False, None
        return True, default
    try:
        f = float(str(v).replace(",", "").strip()) if isinstance(v, str) else float(v)
    except (TypeError, ValueError):
        issues.append(Issue(code="NOT_A_NUMBER", severity="error", message=f"{label} '{v}' is not a number", ref=ref))
        return False, None
    if f != f:  # NaN
        issues.append(Issue(code="NOT_A_NUMBER", severity="error", message=f"{label} is not a number", ref=ref))
        return False, None
    if positive and (f < 0 or (f == 0 and not allow_zero)):
        issues.append(Issue(code="NON_POSITIVE_VALUE", severity="error",
                            message=f"{label} must be {'>= 0' if allow_zero else '> 0'} (got {v})", ref=ref))
        return False, None
    if integer:
        if abs(f - round(f)) > 1e-9:
            issues.append(Issue(code="NOT_AN_INTEGER", severity="error", message=f"{label} must be a whole number",
                                ref=ref))
            return False, None
        f = int(round(f))
    return True, f


def _text(d: Dict[str, Any], names: List[str]) -> Optional[str]:
    for nm in names:
        v = d.get(_key(nm))
        if v is not None and str(v).strip() != "":
            return str(v).strip()
    return None


def _bool(d: Dict[str, Any], names: List[str], default: bool, issues: List[Issue], ref: str, label: str) -> bool:
    v = _text(d, names)
    if v is None:
        return default
    s = v.strip().lower()
    if s in ("y", "yes", "true", "1", "1.0"):
        return True
    if s in ("n", "no", "false", "0", "0.0"):
        return False
    issues.append(Issue(code="BAD_BOOLEAN", severity="warning",
                        message=f"{label} '{v}' not understood; default {'Y' if default else 'N'} used", ref=ref))
    return default


def _settings(wb, issues: List[Issue]) -> Dict[str, Any]:
    ws = _find_ws(wb, "Settings")
    out: Dict[str, Any] = {}
    if ws is None:
        return out
    for n, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if n == 1 or not row or row[0] is None:
            continue
        name = _key(row[0])
        val = row[1] if len(row) > 1 else None
        if val is None or str(val).strip() == "":
            continue
        out[name] = val
    return out


def _settings_num(raw: Dict[str, Any], name: str, issues: List[Issue]) -> Optional[float]:
    v = raw.get(_key(name))
    if v is None:
        return None
    try:
        f = float(v)
    except (TypeError, ValueError):
        issues.append(Issue(code="BAD_SETTING", severity="warning", message=f"Setting {name} '{v}' ignored",
                            ref=f"Settings!{name}"))
        return None
    if f < 0:
        issues.append(Issue(code="BAD_SETTING", severity="warning", message=f"Setting {name} negative, ignored",
                            ref=f"Settings!{name}"))
        return None
    return f


def _settings_bool(raw: Dict[str, Any], name: str) -> Optional[bool]:
    v = raw.get(_key(name))
    if v is None:
        return None
    s = str(v).strip().lower()
    if s in ("y", "yes", "true", "1", "1.0"):
        return True
    if s in ("n", "no", "false", "0", "0.0"):
        return False
    return None


def _open(data: bytes, issues: List[Issue]):
    try:
        return load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    except Exception as exc:  # noqa: BLE001 - corrupt upload must not raise
        issues.append(Issue(code="WORKBOOK_UNREADABLE", severity="error", message=f"Cannot read workbook: {exc}"))
        return None


# ---------------------------------------------------------------------------
# public parsers
# ---------------------------------------------------------------------------
def parse_nesting_workbook(data: bytes) -> Tuple[List[Piece2D], List[Sheet2D], Dict[str, Any], List[Issue]]:
    issues: List[Issue] = []
    pieces: List[Piece2D] = []
    sheets: List[Sheet2D] = []
    wb = _open(data, issues)
    if wb is None:
        return pieces, sheets, {}, issues
    ws = _find_ws(wb, "Pieces") or wb.worksheets[0]
    for n, d in _rows(ws, "Pieces", issues):
        ref = f"Pieces!row {n}"
        mark = _text(d, ["Mark", "Id", "Piece"]) or f"R{n}"
        ok1, w = _num(d, ["Width_mm", "Width", "W"], required=True, positive=True, integer=False, issues=issues,
                      ref=ref, label="Width_mm")
        ok2, h = _num(d, ["Height_mm", "Height", "H"], required=True, positive=True, integer=False, issues=issues,
                      ref=ref, label="Height_mm")
        ok3, q = _num(d, ["Qty", "Quantity"], required=False, positive=True, integer=True, issues=issues, ref=ref,
                      label="Qty", default=1)
        ok4, t = _num(d, ["Thickness_mm", "Thickness"], required=False, positive=True, integer=False, issues=issues,
                      ref=ref, label="Thickness_mm")
        if not (ok1 and ok2 and ok3 and ok4):
            continue
        rot = _bool(d, ["CanRotate", "Rotate"], True, issues, ref, "CanRotate")
        pieces.append(Piece2D(id=mark, label=mark, width=w, height=h, quantity=q,
                              material=_text(d, ["Material"]), thickness=t, can_rotate=rot))
    ws = _find_ws(wb, "Stock") or _find_ws(wb, "Sheets")
    if ws is None:
        issues.append(Issue(code="NO_STOCK_SHEET", severity="warning", message="No 'Stock' worksheet found"))
    else:
        for n, d in _rows(ws, "Stock", issues):
            ref = f"Stock!row {n}"
            sid = _text(d, ["Id", "Sheet", "Mark"]) or f"S{n}"
            ok1, w = _num(d, ["Width_mm", "Width"], required=True, positive=True, integer=False, issues=issues,
                          ref=ref, label="Width_mm")
            ok2, h = _num(d, ["Height_mm", "Height"], required=True, positive=True, integer=False, issues=issues,
                          ref=ref, label="Height_mm")
            ok3, q = _num(d, ["Qty", "Quantity"], required=False, positive=True, integer=True, issues=issues,
                          ref=ref, label="Qty", default=None, allow_zero=True)
            ok4, t = _num(d, ["Thickness_mm", "Thickness"], required=False, positive=True, integer=False,
                          issues=issues, ref=ref, label="Thickness_mm")
            ok5, c = _num(d, ["CostPerSheet", "Cost"], required=False, positive=True, integer=False, issues=issues,
                          ref=ref, label="CostPerSheet", allow_zero=True)
            if not (ok1 and ok2 and ok3 and ok4 and ok5):
                continue
            sheets.append(Sheet2D(id=sid, label=sid, width=w, height=h, quantity=q, material=_text(d, ["Material"]),
                                  thickness=t, is_residual=_bool(d, ["IsResidual", "Residual"], False, issues, ref,
                                                                 "IsResidual"),
                                  cost_per_sheet=c))
    raw = _settings(wb, issues)
    settings: Dict[str, Any] = {}
    for name, field in (("Kerf_mm", "kerf"), ("MinResidualWidth_mm", "min_residual_width"),
                        ("MinResidualHeight_mm", "min_residual_height")):
        v = _settings_num(raw, name, issues)
        if v is not None:
            settings[field] = v
    trim = {}
    for side in ("left", "right", "top", "bottom"):
        v = _settings_num(raw, f"Trim{side.capitalize()}_mm", issues)
        if v is not None:
            trim[side] = v
    if trim:
        settings["edge_trim"] = trim
    for name, field in (("Guillotine", "guillotine"), ("PreferResiduals", "prefer_residuals")):
        b = _settings_bool(raw, name)
        if b is not None:
            settings[field] = b
    if not pieces:
        issues.append(Issue(code="NO_PIECES", severity="warning", message="No valid piece rows found"))
    return pieces, sheets, settings, issues


def parse_cutting_workbook(data: bytes) -> Tuple[List[Piece1D], List[Stock1D], Dict[str, Any], List[Issue]]:
    issues: List[Issue] = []
    pieces: List[Piece1D] = []
    stocks: List[Stock1D] = []
    wb = _open(data, issues)
    if wb is None:
        return pieces, stocks, {}, issues
    ws = _find_ws(wb, "Pieces") or wb.worksheets[0]
    for n, d in _rows(ws, "Pieces", issues):
        ref = f"Pieces!row {n}"
        mark = _text(d, ["Mark", "Id", "Piece"]) or f"R{n}"
        ok1, L = _num(d, ["Length_mm", "Length", "L"], required=True, positive=True, integer=False, issues=issues,
                      ref=ref, label="Length_mm")
        ok2, q = _num(d, ["Qty", "Quantity"], required=False, positive=True, integer=True, issues=issues, ref=ref,
                      label="Qty", default=1)
        ok3, a1 = _num(d, ["AngleStart"], required=False, positive=True, integer=False, issues=issues, ref=ref,
                       label="AngleStart", default=90.0)
        ok4, a2 = _num(d, ["AngleEnd"], required=False, positive=True, integer=False, issues=issues, ref=ref,
                       label="AngleEnd", default=90.0)
        if not (ok1 and ok2 and ok3 and ok4):
            continue
        pieces.append(Piece1D(id=mark, label=_text(d, ["Label", "Description"]) or mark, length=L, quantity=q,
                              angle_start=a1, angle_end=a2))
    ws = _find_ws(wb, "Stock")
    if ws is None:
        issues.append(Issue(code="NO_STOCK_SHEET", severity="warning", message="No 'Stock' worksheet found"))
    else:
        for n, d in _rows(ws, "Stock", issues):
            ref = f"Stock!row {n}"
            sid = _text(d, ["Id", "Mark"]) or f"S{n}"
            ok1, L = _num(d, ["Length_mm", "Length"], required=True, positive=True, integer=False, issues=issues,
                          ref=ref, label="Length_mm")
            ok2, q = _num(d, ["Qty", "Quantity"], required=False, positive=True, integer=True, issues=issues,
                          ref=ref, label="Qty", default=None, allow_zero=True)
            ok3, c = _num(d, ["CostPerBar", "Cost"], required=False, positive=True, integer=False, issues=issues,
                          ref=ref, label="CostPerBar", allow_zero=True)
            if not (ok1 and ok2 and ok3):
                continue
            stocks.append(Stock1D(id=sid, label=sid, length=L, quantity=q, cost_per_bar=c,
                                  is_offcut=_bool(d, ["IsOffcut", "Offcut"], False, issues, ref, "IsOffcut")))
    raw = _settings(wb, issues)
    settings: Dict[str, Any] = {}
    for name, field in (("Kerf_mm", "kerf"), ("TrimStart_mm", "trim_start"), ("TrimEnd_mm", "trim_end"),
                        ("MinReusableOffcut_mm", "min_reusable_offcut")):
        v = _settings_num(raw, name, issues)
        if v is not None:
            settings[field] = v
    b = _settings_bool(raw, "PreferOffcuts")
    if b is not None:
        settings["prefer_offcuts"] = b
    mat = raw.get(_key("Material"))
    if mat is not None and str(mat).strip():
        settings["material"] = str(mat).strip()
    if not pieces:
        issues.append(Issue(code="NO_PIECES", severity="warning", message="No valid piece rows found"))
    return pieces, stocks, settings, issues
