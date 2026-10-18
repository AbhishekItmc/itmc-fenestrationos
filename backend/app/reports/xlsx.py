"""Excel reports (openpyxl). First sheet "Info" carries the full metadata; every sheet starts with the banner
row(s) when provisional/invalid. Workshop sheets are flat tables (one row per cut) for saw operators."""
from __future__ import annotations

import io
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..engines.models import CalcResult, Cutting1DResult, Nesting2DResult
from .common import ReportMeta

RED = PatternFill("solid", fgColor="B71C1C")
HEAD = PatternFill("solid", fgColor="E3E7EC")


class _Book:
    def __init__(self, meta: ReportMeta):
        self.meta = meta
        self.wb = Workbook()
        ws = self.wb.active
        ws.title = "Info"
        r = self._banner(ws, 1, 2)
        ws.cell(r, 1, meta.title).font = Font(bold=True, size=14)
        r += 1
        for k, v in meta.rows():
            ws.cell(r, 1, k).font = Font(bold=True)
            ws.cell(r, 2, v)
            r += 1
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 90

    def _banner(self, ws, row: int, span: int) -> int:
        for b in self.meta.banners:
            c = ws.cell(row, 1, b)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = RED
            for col in range(2, max(span, 2) + 1):
                ws.cell(row, col).fill = RED
            row += 1
        return row

    def sheet(self, title: str, headers: List[str], rows: List[List[Any]],
              widths: Optional[List[int]] = None) -> None:
        ws = self.wb.create_sheet(title)
        r = self._banner(ws, 1, len(headers))
        ws.cell(r, 1, f"Units: {self.meta.units} | Status: {self.meta.validation_status} | "
                      f"{self.meta.reference}").font = Font(italic=True, size=8)
        r += 1
        for j, h in enumerate(headers, 1):
            c = ws.cell(r, j, h)
            c.font = Font(bold=True)
            c.fill = HEAD
            c.alignment = Alignment(wrap_text=True)
        ws.freeze_panes = ws.cell(r + 1, 1)
        for row in rows:
            r += 1
            for j, v in enumerate(row, 1):
                ws.cell(r, j, v)
        for j, h in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(j)].width = (widths[j - 1] if widths and j - 1 < len(widths)
                                                                else max(10, min(40, len(h) + 4)))

    def issues(self, issues: List[Dict[str, Any]], title: str = "Issues") -> None:
        self.sheet(title, ["Severity", "Code", "Message", "Ref"],
                   [[i.get("severity"), i.get("code"), i.get("message"), i.get("ref")] for i in issues],
                   [10, 30, 90, 25])

    def bytes(self) -> bytes:
        buf = io.BytesIO()
        self.wb.save(buf)
        return buf.getvalue()


def cutting_xlsx(result: Dict[str, Any], meta: ReportMeta, name: str) -> bytes:
    r = Cutting1DResult.model_validate(result)
    rq, s = r.request, r.summary
    b = _Book(meta)
    summ = [["Job", name], ["Material / profile", rq.material], ["Optimiser status", r.status],
            ["Method", r.method], ["Bars used", s.bars_used], ["Total stock length (mm)", s.total_stock_length],
            ["Total piece length (mm)", s.total_piece_length], ["Kerf loss (mm)", s.total_kerf_loss],
            ["Trim loss (mm)", s.total_trim_loss], ["Reusable offcut length (mm)", s.reusable_offcut_length],
            ["Waste length (mm)", s.waste_length], ["Utilisation", round(s.utilization, 6)],
            ["Kerf (mm)", rq.kerf], ["Trim start (mm)", rq.trim_start], ["Trim end (mm)", rq.trim_end],
            ["Min reusable offcut (mm)", rq.min_reusable_offcut],
            ["Validation", f"{r.validation.validator}: {'PASSED' if r.validation.passed else 'FAILED'}"]]
    if r.proof:
        summ += [["Lower bound bars", r.proof.lower_bound_bars], ["Proven optimal", r.proof.is_optimal]]
    b.sheet("Summary", ["Item", "Value"], summ, [30, 60])
    # workshop sheet: one row per cut per physical bar
    rows, bar_no = [], 0
    for p in r.patterns:
        for _ in range(p.count):
            bar_no += 1
            for i, c in enumerate(p.cuts, 1):
                rows.append([bar_no, p.pattern_id, p.stock_id, p.stock_length, "yes" if p.is_offcut_stock else "no",
                             i, c.piece_id, c.label, c.length, c.angle_start, c.angle_end, c.offset,
                             round(c.offset + c.length, 3), p.remainder, "offcut" if p.offcut_reusable else "waste"])
    b.sheet("Bars", ["Bar no", "Pattern", "Stock id", "Stock length", "Offcut stock", "Cut order", "Piece", "Label",
                     "Length", "Angle start", "Angle end", "Offset start", "Offset end", "Remainder",
                     "Remainder type"], rows, [8, 9, 18, 12, 10, 9, 12, 30, 10, 10, 10, 12, 12, 11, 13])
    rows = []
    for p in r.patterns:
        for i, c in enumerate(p.cuts, 1):
            rows.append([p.pattern_id, p.count, p.stock_id, p.stock_length, i, c.piece_id, c.label, c.length,
                         c.angle_start, c.angle_end, c.offset])
    b.sheet("Patterns", ["Pattern", "Bar count", "Stock id", "Stock length", "Cut order", "Piece", "Label", "Length",
                         "Angle start", "Angle end", "Offset"], rows)
    placed: Dict[str, int] = {}
    for p in r.patterns:
        for c in p.cuts:
            placed[c.piece_id] = placed.get(c.piece_id, 0) + p.count
    b.sheet("CutList", ["Piece", "Label", "Length", "Angle start", "Angle end", "Qty required", "Qty placed"],
            [[pc.id, pc.label, pc.length, pc.angle_start, pc.angle_end, pc.quantity, placed.get(pc.id, 0)]
             for pc in rq.pieces], [12, 40, 10, 10, 10, 12, 12])
    b.sheet("Offcuts", ["Pattern", "Bars", "Remainder (mm)", "Reusable", "Material"],
            [[p.pattern_id, p.count, p.remainder, "yes" if p.offcut_reusable else "no", rq.material]
             for p in r.patterns])
    b.issues([i.model_dump() for i in r.issues] + [e.model_dump() for e in r.validation.errors])
    return b.bytes()


def nesting_xlsx(result: Dict[str, Any], meta: ReportMeta, name: str) -> bytes:
    r = Nesting2DResult.model_validate(result)
    rq, s = r.request, r.summary
    b = _Book(meta)
    b.sheet("Summary", ["Item", "Value"], [
        ["Job", name], ["Optimiser status", r.status], ["Method", r.method], ["Sheets used", s.sheets_used],
        ["Total sheet area (m²)", s.total_sheet_area_m2], ["Total piece area (m²)", s.total_piece_area_m2],
        ["Waste area (m²)", s.waste_area_m2], ["Reusable residual area (m²)", s.reusable_residual_area_m2],
        ["Utilisation", round(s.utilization, 6)], ["Kerf (mm)", rq.kerf], ["Guillotine", rq.guillotine],
        ["Validation", f"{r.validation.validator}: {'PASSED' if r.validation.passed else 'FAILED'}"]], [30, 60])
    rows, n = [], 0
    for lay in r.layouts:
        for _ in range(lay.count):
            n += 1
            for i, p in enumerate(lay.placements, 1):
                rows.append([n, lay.layout_id, lay.sheet_id, lay.width, lay.height, i, p.piece_id, p.label,
                             p.x, p.y, p.width, p.height, "yes" if p.rotated else "no"])
    b.sheet("Layouts", ["Sheet no", "Layout", "Sheet id", "Sheet W", "Sheet H", "Seq", "Piece", "Label", "x", "y",
                        "W placed", "H placed", "Rotated"], rows)
    rows = []
    for lay in r.layouts:
        for res in lay.residuals:
            rows.append([lay.layout_id, lay.count, res.x, res.y, res.width, res.height,
                         "yes" if res.reusable else "no"])
    b.sheet("Residuals", ["Layout", "Count", "x", "y", "W", "H", "Reusable"], rows)
    placed: Dict[str, int] = {}
    for lay in r.layouts:
        for p in lay.placements:
            placed[p.piece_id] = placed.get(p.piece_id, 0) + lay.count
    b.sheet("Pieces", ["Piece", "Label", "Width", "Height", "Qty", "Placed", "Material", "Thickness", "Can rotate"],
            [[p.id, p.label, p.width, p.height, p.quantity, placed.get(p.id, 0), p.material, p.thickness,
              p.can_rotate] for p in rq.pieces])
    b.issues([i.model_dump() for i in r.issues] + [e.model_dump() for e in r.validation.errors])
    return b.bytes()


def calc_xlsx(result: Dict[str, Any], meta: ReportMeta, name: str) -> bytes:
    r = CalcResult.model_validate(result)
    b = _Book(meta)
    b.sheet("ProfileCuts", ["Mark", "Role", "Profile id", "Profile code", "Description", "Length (mm)",
                            "Angle start", "Angle end", "Qty", "Position", "Formula", "Rules"],
            [[c.mark, c.role, c.profile_id, c.profile_code, c.description, c.length, c.angle_start, c.angle_end,
              c.quantity, c.position, c.formula, "; ".join(c.rule_ids)] for c in r.profile_cuts])
    b.sheet("Glass", ["Mark", "Field", "Fill type", "Glass type", "Panel type", "Description", "Width (mm)",
                      "Height (mm)", "Area (m²)", "Qty", "Formula"],
            [[g.mark, g.field_key, g.fill_type, g.glass_type_id, g.panel_type_id, g.description, g.width, g.height,
              g.area_m2, g.quantity, g.formula] for g in r.glass_items])
    b.sheet("BOM", ["Type", "Item id", "Code", "Description", "Unit", "Quantity", "Rules", "Demo data"],
            [[x.item_type, x.item_id, x.code, x.description, x.unit, x.quantity, "; ".join(x.rule_ids), x.demo_data]
             for x in r.bom])
    b.sheet("Totals", ["Key", "Value"], [[k, v] for k, v in r.totals.items()], [40, 20])
    b.issues([i.model_dump() for i in r.issues] + [e.model_dump() for e in r.validation.errors])
    return b.bytes()


def quote_xlsx(q: Dict[str, Any], meta: ReportMeta) -> bytes:
    b = _Book(meta)
    rows = [[ln["line_no"], ln["item_type"], ln["code"], ln["description"], ln["unit"], ln["quantity"],
             ln["unit_price"], ln["amount"]] for ln in q["lines"]]
    rows += [[], ["", "", "", "Subtotal", "", "", "", q["subtotal"]],
             ["", "", "", f"Margin {q['margin_pct']:g} %", "", "", "", q["margin_amount"]],
             ["", "", "", "TOTAL", "", "", "", q["total"]],
             ["", "", "", "Demo prices — not a commercial offer"]]
    b.sheet("Lines", ["#", "Type", "Code", "Description", "Unit", "Qty", f"Unit price ({q['currency']})",
                      f"Amount ({q['currency']})"], rows, [5, 14, 20, 45, 6, 12, 16, 16])
    return b.bytes()


def purchase_order_xlsx(po: Dict[str, Any], meta: ReportMeta) -> bytes:
    b = _Book(meta)
    cur = po["currency"]
    rows = [[ln.get("line_no"), ln["item_type"], ln["code"], ln.get("description"), ln["unit"],
             ln.get("qty_required"), ln.get("qty_in_stock"), ln.get("qty_to_order"), ln.get("unit_price"),
             ln.get("amount"), ln.get("stock_length"), ln.get("note")] for ln in po["lines"]]
    rows += [[], ["", "", "", "TOTAL (excl. taxes)", "", "", "", "", "", po["total"]],
             ["", "", "", "Demo prices — not a commercial offer"]]
    b.sheet("Lines", ["#", "Type", "Code", "Description", "Unit", "Qty required", "Qty in stock", "Qty to order",
                      f"Unit price ({cur})", f"Amount ({cur})", "Bar length (mm)", "Note"], rows,
            [5, 12, 20, 45, 6, 12, 12, 12, 14, 14, 12, 45])
    return b.bytes()
