import io

from openpyxl import load_workbook

from helpers import assert_error

CUT_REQ = {"material": "TEST-PROFILE", "kerf": 5.0, "trim_start": 10.0, "trim_end": 20.0,
           "min_reusable_offcut": 500.0, "stocks": [{"id": "S6000", "length": 6000}],
           "pieces": [{"id": "A", "length": 1450, "quantity": 7}, {"id": "B", "length": 2300, "quantity": 3},
                      {"id": "C", "length": 610, "quantity": 5, "angle_start": 45}]}

NEST_REQ = {"kerf": 4.0, "sheets": [{"id": "G1", "width": 3210, "height": 2250, "material": "glass", "thickness": 6}],
            "pieces": [{"id": "P1", "width": 1000, "height": 800, "quantity": 6, "material": "glass", "thickness": 6},
                       {"id": "P2", "width": 1500, "height": 600, "quantity": 3, "material": "glass", "thickness": 6,
                        "can_rotate": False}],
            "guillotine": True}


def _check_kerf(result, kerf):
    rq = result["request"]
    for p in result["patterns"]:
        off = rq["trim_start"]
        for c in p["cuts"]:
            assert abs(c["offset"] - off) < 1e-6
            off += c["length"] + kerf
        assert abs(p["used_length"] - off) < 1e-6
        assert abs(p["remainder"] - (p["stock_length"] - rq["trim_end"] - off)) < 1e-6
        assert p["remainder"] >= -1e-6


def test_cutting_optimize_stateless(client, auth):
    r = client.post("/api/v1/cutting/1d/optimize", headers=auth["viewer"], json=CUT_REQ)
    assert r.status_code == 200
    res = r.json()
    assert res["validation"]["passed"] is True
    assert res["validation_status"] in ("FEASIBLE_HEURISTIC", "OPTIMAL_PROVEN")
    assert not res["unplaced"]
    _check_kerf(res, 5.0)
    placed = sum(len(p["cuts"]) * p["count"] for p in res["patterns"])
    assert placed == 15
    assert any(i["code"] == "MITRE_ALLOWANCE_NOT_APPLIED" for i in res["issues"])


def test_cutting_optimize_bad_request(client, auth):
    bad = {**CUT_REQ, "kerf": -1}
    assert_error(client.post("/api/v1/cutting/1d/optimize", headers=auth["viewer"], json=bad), 422)
    too_long = {**CUT_REQ, "pieces": [{"id": "X", "length": 7000, "quantity": 1}]}
    res = client.post("/api/v1/cutting/1d/optimize", headers=auth["viewer"], json=too_long).json()
    assert res["unplaced"] and any(i["code"] == "PIECE_TOO_LONG" for i in res["issues"])


def test_nesting_optimize_stateless(client, auth):
    r = client.post("/api/v1/nesting/2d/optimize", headers=auth["viewer"], json=NEST_REQ)
    assert r.status_code == 200
    res = r.json()
    assert res["validation"]["passed"] is True and res["validation_status"] != "INVALID"
    assert not res["unplaced"]
    for lay in res["layouts"]:
        pl = lay["placements"]
        for i in range(len(pl)):
            for j in range(i + 1, len(pl)):
                a, b = pl[i], pl[j]
                gap_x = max(b["x"] - (a["x"] + a["width"]), a["x"] - (b["x"] + b["width"]))
                gap_y = max(b["y"] - (a["y"] + a["height"]), a["y"] - (b["y"] + b["height"]))
                assert max(gap_x, gap_y) >= 4.0 - 1e-6     # kerf honoured
        for p in pl:
            if p["piece_id"] == "P2":
                assert not p["rotated"]


def test_cutting_jobs_persist(client, auth, project):
    p, _ = project
    r = client.post("/api/v1/cutting/jobs", headers=auth["engineer"],
                    json={"project_id": p["id"], "name": "Job A", "request": CUT_REQ})
    assert r.status_code == 201, r.text
    job = r.json()
    assert job["status"] == "completed" and job["result"]["validation"]["passed"]
    assert job["summary"]["bars_used"] == job["result"]["summary"]["bars_used"]
    got = client.get(f"/api/v1/cutting/jobs/{job['id']}", headers=auth["viewer"]).json()
    assert got["result"]["patterns"] == job["result"]["patterns"]
    lst = client.get(f"/api/v1/cutting/jobs?project_id={p['id']}", headers=auth["viewer"]).json()
    assert [j["id"] for j in lst] == [job["id"]] and "result" not in lst[0]
    # draft job then optimise
    d = client.post("/api/v1/cutting/jobs", headers=auth["engineer"],
                    json={"name": "Draft", "request": CUT_REQ, "optimize": False}).json()
    assert d["status"] == "draft" and d["result"] is None
    assert_error(client.get(f"/api/v1/reports/cutting/{d['id']}.pdf", headers=auth["viewer"]), 400)
    o = client.post(f"/api/v1/cutting/jobs/{d['id']}/optimize", headers=auth["engineer"]).json()
    assert o["status"] == "completed"
    assert client.delete(f"/api/v1/cutting/jobs/{d['id']}", headers=auth["engineer"]).status_code == 204
    assert_error(client.post("/api/v1/cutting/jobs", headers=auth["engineer"],
                             json={"project_id": 999999, "name": "x", "request": CUT_REQ}), 400, "PROJECT_NOT_FOUND")


def test_nesting_jobs_persist(client, auth):
    r = client.post("/api/v1/nesting/jobs", headers=auth["engineer"], json={"name": "Glass A", "request": NEST_REQ})
    assert r.status_code == 201
    job = r.json()
    assert job["status"] == "completed" and job["validation_status"] != "INVALID"
    assert client.get(f"/api/v1/nesting/jobs/{job['id']}", headers=auth["viewer"]).json()["id"] == job["id"]
    assert any(j["id"] == job["id"] for j in client.get("/api/v1/nesting/jobs", headers=auth["viewer"]).json())


def test_templates_download(client, auth):
    for path in ("/api/v1/nesting/template.xlsx", "/api/v1/cutting/template.xlsx"):
        r = client.get(path, headers=auth["viewer"])
        assert r.status_code == 200 and r.content[:2] == b"PK"
        load_workbook(io.BytesIO(r.content))


def _fill_nesting_template(client, auth) -> bytes:
    wb = load_workbook(io.BytesIO(client.get("/api/v1/nesting/template.xlsx", headers=auth["viewer"]).content))
    ws = wb["Pieces"]
    for row in list(ws.iter_rows(min_row=2)):
        for c in row:
            c.value = None
    rows = [["W1", 1200, 900, 4, "glass", 6, "Y"], ["W2", 700, 2100, 2, "glass", 6, "N"],
            ["BAD", -5, 100, 1, "glass", 6, "Y"]]
    for i, r in enumerate(rows, start=2):
        for j, v in enumerate(r, start=1):
            ws.cell(i, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_nesting_excel_import(client, auth):
    data = _fill_nesting_template(client, auth)
    r = client.post("/api/v1/nesting/import-excel", headers=auth["viewer"],
                    files={"file": ("pieces.xlsx", data, "application/vnd.openxmlformats-officedocument."
                                                         "spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    b = r.json()
    ids = {p["id"]: p for p in b["pieces"]}
    assert ids["W1"]["quantity"] == 4 and ids["W2"]["can_rotate"] is False
    assert "BAD" not in ids
    assert any(i["severity"] == "error" for i in b["issues"])
    # optimise what was imported
    assert {sh["id"] for sh in b["sheets"]} == {"SHEET-3210x2250", "ACP-1500x4000"}
    req = {"pieces": b["pieces"], **b["settings"],
           "sheets": [{"id": "S", "width": 3210, "height": 2250, "material": "glass", "thickness": 6}]}
    res = client.post("/api/v1/nesting/2d/optimize", headers=auth["viewer"], json=req).json()
    assert res["validation"]["passed"] and not res["unplaced"]


def test_excel_import_rejects_garbage(client, auth):
    r = client.post("/api/v1/nesting/import-excel", headers=auth["viewer"],
                    files={"file": ("x.txt", b"hello world", "text/plain")})
    assert_error(r, 415, "UNSUPPORTED_FILE")
    r = client.post("/api/v1/cutting/import-excel", headers=auth["viewer"],
                    files={"file": ("x.xlsx", b"PK\x03\x04garbage", "application/octet-stream")})
    assert r.status_code == 200 and any(i["code"] == "WORKBOOK_UNREADABLE" for i in r.json()["issues"])
    assert_error(client.post("/api/v1/nesting/import-excel", headers=auth["viewer"]), 422)


def test_cutting_excel_import_template_roundtrip(client, auth):
    data = client.get("/api/v1/cutting/template.xlsx", headers=auth["viewer"]).content
    r = client.post("/api/v1/cutting/import-excel", headers=auth["viewer"], files={"file": ("t.xlsx", data)})
    assert r.status_code == 200
    b = r.json()
    assert "pieces" in b and "stocks" in b and isinstance(b["issues"], list)
