import io
import math

from openpyxl import load_workbook
from pypdf import PdfReader

from helpers import assert_error
from test_optimisation import CUT_REQ, NEST_REQ

BANNER = "PROVISIONAL — DEMO RULES/DATA — NOT FOR PRODUCTION"


def pdf_text(data: bytes) -> str:
    assert data[:5] == b"%PDF-"
    return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)


def xlsx_info(data: bytes):
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames[0] == "Info"
    info = {r[0]: r[1] for r in wb["Info"].iter_rows(values_only=True) if r[0]}
    return wb, info


def _calc_run(client, auth, project):
    p, e = project
    return client.post(f"/api/v1/elevations/{e['id']}/calculate", headers=auth["engineer"]).json()


def test_calc_reports_provisional(client, auth, project):
    run = _calc_run(client, auth, project)
    r = client.get(f"/api/v1/reports/calc/{run['id']}.pdf", headers=auth["viewer"])
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.headers["X-Provisional"] == "true"
    txt = pdf_text(r.content)
    assert "NOT FOR PRODUCTION" in txt and "PROVISIONAL" in txt
    for needle in ("PROVISIONAL_DEMO_RULES", "calc/", run["catalogue_version"], "rs-cw50-demo", "Units", "mm",
                   "Generated at", "Profile cut list", "Bill of materials"):
        assert needle in txt, needle
    r = client.get(f"/api/v1/reports/calc/{run['id']}.xlsx", headers=auth["viewer"])
    assert r.status_code == 200
    wb, info = xlsx_info(r.content)
    assert {"ProfileCuts", "Glass", "BOM", "Issues"} <= set(wb.sheetnames)
    assert info["Validation status"] == "PROVISIONAL_DEMO_RULES"
    assert info["Catalogue version"] == run["catalogue_version"] and info["Units"].startswith("mm")
    assert wb["Info"]["A1"].value == BANNER and wb["BOM"]["A1"].value == BANNER


def test_cutting_reports(client, auth):
    job = client.post("/api/v1/cutting/jobs", headers=auth["engineer"], json={"name": "Rep", "request": CUT_REQ}).json()
    r = client.get(f"/api/v1/reports/cutting/{job['id']}.pdf", headers=auth["viewer"])
    assert r.status_code == 200
    txt = pdf_text(r.content)
    assert "cutting1d/" in txt and job["validation_status"] in txt and "Cutting patterns" in txt
    assert BANNER not in txt          # user-supplied (non-demo) job -> no demo banner
    r = client.get(f"/api/v1/reports/cutting/{job['id']}.xlsx", headers=auth["viewer"])
    wb, info = xlsx_info(r.content)
    assert info["Engine version"].startswith("cutting1d/")
    bars = list(wb["Bars"].iter_rows(values_only=True))
    hdr_idx = next(i for i, row in enumerate(bars) if row[0] == "Bar no")
    hdr = bars[hdr_idx]
    assert {"Pattern", "Cut order", "Length", "Angle start", "Angle end", "Offset start"} <= set(hdr)
    data = bars[hdr_idx + 1:]
    assert len(data) == 15   # one row per piece cut
    assert max(row[0] for row in data) == job["result"]["summary"]["bars_used"]
    assert_error(client.get("/api/v1/reports/cutting/999999.pdf", headers=auth["viewer"]), 404)
    assert client.get(f"/api/v1/reports/cutting/{job['id']}.doc", headers=auth["viewer"]).status_code == 422


def test_nesting_reports(client, auth):
    job = client.post("/api/v1/nesting/jobs", headers=auth["engineer"], json={"name": "Rep", "request": NEST_REQ}).json()
    r = client.get(f"/api/v1/reports/nesting/{job['id']}.pdf", headers=auth["viewer"])
    txt = pdf_text(r.content)
    assert "nesting2d/" in txt and "Sheet layouts" in txt and "P1" in txt
    wb, info = xlsx_info(client.get(f"/api/v1/reports/nesting/{job['id']}.xlsx", headers=auth["viewer"]).content)
    assert {"Layouts", "Residuals", "Pieces"} <= set(wb.sheetnames)


def test_demo_cutting_job_from_calc_has_banner(client, auth, project):
    run = _calc_run(client, auth, project)
    items = client.post(f"/api/v1/cutting/from-calc/{run['id']}?optimize=true", headers=auth["engineer"]).json()
    jid = items[0]["job_id"]
    txt = pdf_text(client.get(f"/api/v1/reports/cutting/{jid}.pdf", headers=auth["viewer"]).content)
    assert "NOT FOR PRODUCTION" in txt and run["catalogue_version"] in txt
    wb, info = xlsx_info(client.get(f"/api/v1/reports/cutting/{jid}.xlsx", headers=auth["viewer"]).content)
    assert wb["Bars"]["A1"].value == BANNER


def test_quote_totals_and_report(client, auth, project):
    p, e = project
    run = _calc_run(client, auth, project)
    r = client.post(f"/api/v1/projects/{p['id']}/quotes", headers=auth["estimator"],
                    json={"calc_run_ids": [run["id"]], "margin_pct": 12.5, "notes": "Budget"})
    assert r.status_code == 201, r.text
    q = r.json()
    assert q["provisional"] is True and q["validation_status"] == "PROVISIONAL_DEMO_RULES"
    assert "Demo prices — not a commercial offer" in q["notes"]
    import re
    assert re.match(r"^Q-\d{4}-\d{4}$", q["number"])
    cat = client.get("/api/v1/catalogue", headers=auth["viewer"]).json()
    profiles = {x["id"]: x for x in cat["profiles"]}
    glass = {x["id"]: x for x in cat["glass_types"]}
    acc = {x["id"]: x for x in cat["accessories"]}
    colour = next(c for c in cat["colours"] if c["id"] == "col-ral7016")
    expected_profile_amount = 0.0
    for ln in q["lines"]:
        assert math.isclose(ln["amount"], round(ln["quantity"] * ln["unit_price"], 2), abs_tol=0.011) or \
            ln["item_type"] == "colour_surcharge"
        if ln["item_type"] == "profile":
            assert ln["unit_price"] == profiles[ln["item_id"]]["price_per_m"]
            expected_profile_amount += ln["quantity"] * ln["unit_price"]
        elif ln["item_type"] == "glass":
            assert ln["unit_price"] == glass[ln["item_id"]]["price_per_m2"]
        elif ln["item_type"] in ("gasket", "connector", "bracket", "accessory", "sealant"):
            assert ln["unit_price"] == acc[ln["item_id"]]["price"]
    bom = run["result"]["bom"]
    assert len([ln for ln in q["lines"] if ln["item_type"] != "colour_surcharge"]) == len(bom)
    sur = [ln for ln in q["lines"] if ln["item_type"] == "colour_surcharge"]
    if colour["surcharge_pct"]:
        assert len(sur) == 1
        assert math.isclose(sur[0]["amount"], round(expected_profile_amount, 2) * colour["surcharge_pct"] / 100,
                            abs_tol=0.011)
    assert math.isclose(q["subtotal"], round(sum(ln["amount"] for ln in q["lines"]), 2), abs_tol=1e-6)
    assert math.isclose(q["margin_amount"], round(q["subtotal"] * 0.125, 2), abs_tol=1e-6)
    assert math.isclose(q["total"], round(q["subtotal"] + q["margin_amount"], 2), abs_tol=1e-6)
    # numbering is sequential
    q2 = client.post(f"/api/v1/projects/{p['id']}/quotes", headers=auth["admin"],
                     json={"calc_run_ids": [run["id"]]}).json()
    assert int(q2["number"][-4:]) == int(q["number"][-4:]) + 1 and q2["number"][:7] == q["number"][:7]
    assert client.get(f"/api/v1/quotes/{q['id']}", headers=auth["viewer"]).json()["total"] == q["total"]
    assert len(client.get(f"/api/v1/projects/{p['id']}/quotes", headers=auth["viewer"]).json()) == 2
    u = client.put(f"/api/v1/quotes/{q['id']}", headers=auth["estimator"], json={"status": "sent"}).json()
    assert u["status"] == "sent"
    # report
    r = client.get(f"/api/v1/reports/quote/{q['id']}.pdf", headers=auth["viewer"])
    txt = pdf_text(r.content)
    assert q["number"] in txt and "NOT FOR PRODUCTION" in txt and "not a commercial offer" in txt
    assert f"{q['total']:,.2f}" in txt
    wb, info = xlsx_info(client.get(f"/api/v1/reports/quote/{q['id']}.xlsx", headers=auth["viewer"]).content)
    assert info["Validation status"] == "PROVISIONAL_DEMO_RULES"


def test_quote_errors(client, auth, project):
    p, e = project
    assert_error(client.post(f"/api/v1/projects/{p['id']}/quotes", headers=auth["estimator"],
                             json={"calc_run_ids": [999999]}), 404, "CALC_RUN_NOT_FOUND")
    assert_error(client.post(f"/api/v1/projects/{p['id']}/quotes", headers=auth["estimator"],
                             json={"calc_run_ids": []}), 422)
    assert_error(client.post(f"/api/v1/projects/{p['id']}/quotes", headers=auth["estimator"],
                             json={"calc_run_ids": [1], "margin_pct": -3}), 422)
