import io

from openpyxl import Workbook, load_workbook

from app.engines.cutting1d import optimize_1d
from app.engines.excel_io import (build_cutting_template, build_nesting_template, parse_cutting_workbook,
                                  parse_nesting_workbook)
from app.engines.models import Cutting1DRequest, Nesting2DRequest
from app.engines.nesting2d import optimize_2d


def _bytes(wb):
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_nesting_template_round_trip():
    data = build_nesting_template()
    wb = load_workbook(io.BytesIO(data))
    assert [ws.title for ws in wb.worksheets] == ["Pieces", "Stock", "Settings"]
    assert [c.value for c in wb["Pieces"][1]] == ["Mark", "Width_mm", "Height_mm", "Qty", "Material", "Thickness_mm",
                                                  "CanRotate"]
    pieces, sheets, settings, issues = parse_nesting_workbook(data)
    assert not [i for i in issues if i.severity == "error"]
    assert [(p.id, p.width, p.height, p.quantity, p.can_rotate) for p in pieces] == [
        ("G-01", 1200, 2400, 4, True), ("P-01", 800, 600, 2, False)]
    assert sheets[0].quantity is None and sheets[1].quantity == 20 and not sheets[0].is_residual
    assert settings["kerf"] == 3.0 and settings["guillotine"] is True
    r = optimize_2d(Nesting2DRequest(pieces=pieces, sheets=sheets, **settings))
    assert r.validation.passed and not r.unplaced


def test_cutting_template_round_trip():
    data = build_cutting_template()
    pieces, stocks, settings, issues = parse_cutting_workbook(data)
    assert not [i for i in issues if i.severity == "error"]
    assert [(p.id, p.length, p.quantity) for p in pieces] == [("M1", 3600, 6), ("T1", 1450, 12)]
    assert [s.length for s in stocks] == [6000, 6500]
    assert settings["kerf"] == 4.0 and settings["trim_start"] == 10.0
    r = optimize_1d(Cutting1DRequest(pieces=pieces, stocks=stocks, **settings))
    assert r.validation.passed


def test_tolerant_headers_and_row_issues():
    wb = Workbook()
    ws = wb.active
    ws.title = "pieces"
    ws.append([" mark ", "WIDTH MM", "height_mm", "QTY", "can rotate"])
    ws.append(["A", 1000, 500, 2, "no"])
    ws.append(["B", -5, 500, 1, "Y"])        # negative
    ws.append(["C", 1000, None, 1, "Y"])     # missing height
    ws.append(["D", "abc", 100, 1, "Y"])     # not a number
    ws.append(["E", 100, 100, 1.5, "Y"])     # non-integer qty
    ws.append([None, None, None, None, None])  # blank row ignored
    ws.append(["F", "1,200", 300, None, "maybe"])  # thousands separator, default qty, bad bool -> warning
    st = wb.create_sheet("STOCK")
    st.append(["id", "width_mm", "height_mm", "qty", "isresidual"])
    st.append(["S1", 3210, 2250, None, "N"])
    st.append(["S2", 0, 2250, 1, "N"])
    pieces, sheets, settings, issues = parse_nesting_workbook(_bytes(wb))
    assert [p.id for p in pieces] == ["A", "F"]
    assert pieces[0].can_rotate is False and pieces[1].width == 1200 and pieces[1].quantity == 1
    refs = {i.ref for i in issues if i.severity == "error"}
    assert {"Pieces!row 3", "Pieces!row 4", "Pieces!row 5", "Pieces!row 6", "Stock!row 3"} <= refs
    assert any(i.code == "BAD_BOOLEAN" for i in issues)
    assert [s.id for s in sheets] == ["S1"] and settings == {}


def test_garbage_bytes_do_not_raise():
    p, s, st, issues = parse_nesting_workbook(b"not an xlsx")
    assert p == [] and issues[0].code == "WORKBOOK_UNREADABLE"
    p, s, st, issues = parse_cutting_workbook(b"")
    assert p == [] and issues[0].severity == "error"
